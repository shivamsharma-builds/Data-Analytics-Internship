"""Generate a compact PowerPoint summary from pipeline outputs."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUTS = ROOT / "outputs"
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)


def add_title(slide, title: str, subtitle: str = "") -> None:
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.25), Inches(12), Inches(0.4))
        box.text_frame.text = subtitle
        box.text_frame.paragraphs[0].font.size = Pt(14)


def add_bullets(slide, bullets: list[str], left=0.9, top=2.0, width=11.5, height=4.6) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(20 if index == 0 else 17)
        paragraph.space_after = Pt(12)


def main() -> None:
    summary = json.loads((OUTPUTS / "analysis_summary.json").read_text(encoding="utf-8"))
    metrics = summary["metrics"]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    accent = RGBColor(0x00, 0x79, 0xF2)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_title(slide, "IBM HR Attrition Analytics", "From workforce signals to focused retention conversations")
    add_bullets(slide, [
        f"{summary['records']:,} employee records · {summary['attrition_rate']:.1f}% observed attrition",
        "Explainable Logistic Regression baseline with held-out evaluation",
        "Interactive dashboard and prioritized review queue",
    ], top=2.2)

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "What the dataset says")
    add_bullets(slide, [
        f"Overtime attrition: {summary['overtime_attrition_rate']:.1f}% vs {summary['non_overtime_attrition_rate']:.1f}% without overtime",
        f"Highest department rate: {summary['top_department']}",
        f"Highest role rate: {summary['top_role']} — interpret with group size context",
        "No missing values in the downloaded Kaggle export",
    ], top=1.8)

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Model performance", "Stratified 80/20 split · class-balanced Logistic Regression")
    add_bullets(slide, [
        f"Accuracy: {metrics['accuracy']:.1%}",
        f"Precision: {metrics['precision']:.1%}",
        f"Recall: {metrics['recall']:.1%}",
        f"ROC AUC: {metrics['roc_auc']:.3f}",
        f"Confusion matrix: TN {metrics['confusion_matrix'][0][0]} · FP {metrics['confusion_matrix'][0][1]} · "
        f"FN {metrics['confusion_matrix'][1][0]} · TP {metrics['confusion_matrix'][1][1]}",
    ], top=1.6)

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Recommended retention motions")
    add_bullets(slide, [
        "Audit overtime concentration before applying broad incentives.",
        "Add structured early-tenure check-ins and role-clarity prompts.",
        "Use high-score records to start human conversations, never automatic decisions.",
        "Track intervention outcomes and fairness checks over time.",
    ], top=1.8)

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Aptos"
                        if run.font.size is None:
                            run.font.size = Pt(16)
                        run.font.color.rgb = accent if shape == slide.shapes.title else RGBColor(0x20, 0x33, 0x3A)

    path = REPORTS / "IBM_HR_Attrition_Executive_Summary.pptx"
    prs.save(path)
    print(path)


if __name__ == "__main__":
    main()