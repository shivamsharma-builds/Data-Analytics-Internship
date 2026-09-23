"""Generate the Word submission report from pipeline outputs."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUTS = ROOT / "outputs"
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)


def pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def main() -> None:
    summary = json.loads((OUTPUTS / "analysis_summary.json").read_text(encoding="utf-8"))
    metrics = summary["metrics"]
    matrix = metrics["confusion_matrix"]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("IBM HR Employee Attrition Analytics")
    run.bold = True
    run.font.size = Pt(22)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("End-to-end data preparation, exploratory analysis, and attrition prioritization").italic = True

    document.add_paragraph(
        "Dataset reference: Kaggle — IBM HR Employee Attrition Analysis. "
        "This report adapts the supplied project roadmap from e-commerce to workforce retention."
    )

    add_heading(document, "1. Executive Summary")
    document.add_paragraph(
        f"The dataset contains {summary['records']:,} employee records and an observed attrition rate of "
        f"{summary['attrition_rate']:.2f}%. The analysis identifies overtime, early-career tenure, lower "
        "satisfaction signals, limited stock-option coverage, and mobility/commute context as useful "
        "prioritization signals. Overtime is the clearest descriptive split: the observed attrition rate is "
        f"{summary['overtime_attrition_rate']:.1f}% for employees working overtime versus "
        f"{summary['non_overtime_attrition_rate']:.1f}% for those who do not."
    )
    document.add_paragraph(
        "The Logistic Regression model is intentionally explainable and evaluated on a held-out stratified "
        "test set. Its probabilities are used to focus human review and retention conversations; they are "
        "not causal explanations and must not be used as automatic employment decisions."
    )

    add_heading(document, "2. Problem Statement and Objectives")
    document.add_paragraph(
        "Organizations need an evidence-based way to understand where attrition is concentrated and which "
        "employee groups warrant a closer retention conversation. This project addresses that need by "
        "cleaning the IBM HR dataset, quantifying descriptive drivers, training a transparent baseline "
        "classifier, and surfacing a ranked review queue."
    )
    add_bullets(
        document,
        [
            "Standardize the Kaggle export and document data quality.",
            "Compare attrition across departments, roles, overtime, tenure, satisfaction, and mobility context.",
            "Separate model inputs from the target and evaluate on held-out data.",
            "Produce operational risk tiers with a clear human-review guardrail.",
        ],
    )

    add_heading(document, "3. Dataset Overview and Data Dictionary")
    document.add_paragraph(
        f"The source has {summary['records']:,} rows and 35 source fields. There are no missing values in "
        "the downloaded file, but the pipeline still contains explicit missing-value handling for portability."
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Field", "Type", "Role", "Meaning"]):
        cell.text = text
    dictionary = [
        ("Attrition", "Categorical", "Target", "Historical Yes/No attrition label"),
        ("EmployeeNumber", "Identifier", "Key", "Employee identifier"),
        ("Department / JobRole", "Categorical", "Segment", "Organizational context"),
        ("OverTime", "Categorical", "Driver", "Work pattern signal"),
        ("MonthlyIncome", "Numeric", "Feature", "Monthly income amount"),
        ("YearsAtCompany", "Numeric", "Feature", "Tenure at the company"),
        ("JobSatisfaction / EnvironmentSatisfaction / WorkLifeBalance", "Ordinal", "Feature", "1–4 satisfaction inputs"),
        ("JobLevel / StockOptionLevel", "Ordinal", "Feature", "Role level and equity context"),
    ]
    for row in dictionary:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    add_heading(document, "4. Phase 1 — Data Preprocessing")
    add_bullets(
        document,
        [
            "Removed duplicate rows and normalized the UTF-8 BOM in the first header.",
            "Trimmed categorical text and replaced blank categorical values with Unknown.",
            "Coerced numeric features and used medians for any future numeric missingness.",
            "Created AttritionFlag = 1 when Attrition is Yes and 0 otherwise.",
            "Preserved the original source fields and wrote a cleaned CSV to outputs/.",
        ],
    )

    add_heading(document, "5. Phase 2 — Exploratory Analysis")
    document.add_paragraph(
        "Each visualization is read using the required observation → insight → hypothesis → recommendation framework."
    )
    eda = [
        (
            "Work pattern exposure",
            f"Observation: overtime employees show {summary['overtime_attrition_rate']:.1f}% observed attrition versus "
            f"{summary['non_overtime_attrition_rate']:.1f}% without overtime.",
            "Insight: workload and schedule sustainability are important retention lenses.",
            "Hypothesis: sustained overtime may reduce work-life balance or increase burnout risk.",
            "Recommendation: review overtime concentration by role and manager, then test workload-balancing interventions.",
        ),
        (
            "Department comparison",
            f"Observation: {summary['top_department']} has the highest observed department-level attrition rate in this dataset.",
            "Insight: retention actions should be targeted rather than organization-wide by default.",
            "Hypothesis: local role mix, workload, and manager context may explain the difference.",
            "Recommendation: pair department-level comparisons with role and sample-size context before acting.",
        ),
        (
            "Job-role concentration",
            f"Observation: {summary['top_role']} has the highest observed role-level attrition rate.",
            "Insight: smaller role groups can carry disproportionately high risk signals.",
            "Hypothesis: career path clarity, scheduling, or role-specific labor-market pressure may contribute.",
            "Recommendation: conduct qualitative listening sessions before designing role-specific interventions.",
        ),
        (
            "Tenure and early-career context",
            "Observation: the pipeline explicitly separates early tenure from longer-tenured employee profiles.",
            "Insight: onboarding and the first manager relationship are actionable retention moments.",
            "Hypothesis: new employees have less organizational attachment and less role clarity.",
            "Recommendation: add 30/60/90-day check-ins and a six-month internal mobility conversation.",
        ),
        (
            "Satisfaction signals",
            "Observation: job satisfaction, environment satisfaction, and work-life balance are available as interpretable 1–4 inputs.",
            "Insight: employee listening data can complement operational signals such as overtime and commute.",
            "Hypothesis: low satisfaction may amplify the effect of workload and limited progression.",
            "Recommendation: use survey follow-up as a conversation starter, not as a deterministic score.",
        ),
    ]
    for title_text, observation, insight, hypothesis, recommendation in eda:
        document.add_paragraph(title_text, style="Heading 3")
        for label, text in [
            ("Observation", observation),
            ("Insight", insight),
            ("Hypothesis", hypothesis),
            ("Recommendation", recommendation),
        ]:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(text)

    add_heading(document, "6. Phase 3 — Feature Engineering and Leakage Safeguards")
    document.add_paragraph(
        "Unlike the e-commerce RFM example in the supplied roadmap, the IBM HR file is an employee-level "
        "snapshot and has no transaction date or customer entity. Therefore, this implementation uses "
        "employee-level operational and satisfaction features rather than inventing time windows or pseudo-RFM fields."
    )
    document.add_paragraph(
        "The target column Attrition is excluded from model inputs. A stratified 80/20 split keeps the "
        "historical class balance similar in training and test sets. All scaling and one-hot encoding are "
        "fit inside the training pipeline to prevent preprocessing leakage."
    )

    add_heading(document, "7. Phase 4 — Logistic Regression Results")
    result_table = document.add_table(rows=1, cols=2)
    result_table.style = "Table Grid"
    for cell, text in zip(result_table.rows[0].cells, ["Metric", "Test result"]):
        cell.text = text
    for label, value in [
        ("Accuracy", pct(metrics["accuracy"])),
        ("Precision", pct(metrics["precision"])),
        ("Recall", pct(metrics["recall"])),
        ("ROC AUC", f"{metrics['roc_auc']:.3f}"),
        ("Train / test records", f"{metrics['train_records']:,} / {metrics['test_records']:,}"),
    ]:
        cells = result_table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    document.add_paragraph(
        f"Confusion matrix (rows = actual, columns = predicted): TN={matrix[0][0]}, FP={matrix[0][1]}, "
        f"FN={matrix[1][0]}, TP={matrix[1][1]}."
    )
    document.add_paragraph(
        "Recall is especially important for a retention review queue because false negatives represent "
        "historical attrition cases that the model did not surface. Precision still matters because an "
        "overly broad queue can waste HR attention and create poor employee experiences."
    )

    add_heading(document, "8. Phase 5 — Risk Tiers and Dashboard")
    document.add_paragraph(
        f"The generated scores produce {summary['priority_count']:,} High, {summary['medium_count']:,} Medium, "
        f"and {summary['low_count']:,} Low probability records using the supplied thresholds: High >70%, "
        "Medium 40–70%, and Low <40%."
    )
    add_bullets(
        document,
        [
            "High: prioritize a human-led check-in, workload review, and career-path conversation.",
            "Medium: offer targeted listening, manager support, and relevant development resources.",
            "Low: continue standard engagement and monitor aggregate changes rather than individual action.",
        ],
    )
    document.add_paragraph(
        "The dashboard uses the same source CSV and adds filters, chart-level CSV exports, a searchable risk "
        "table, a dark-mode toggle, print/PDF support, and refresh controls with a five-minute minimum interval."
    )

    add_heading(document, "9. Strategic Recommendations")
    add_bullets(
        document,
        [
            "Audit overtime concentration by role, team, and manager before launching broad retention programs.",
            "Strengthen early-tenure onboarding with 30/60/90-day check-ins and role-clarity prompts.",
            "Use department and job-role comparisons to focus qualitative listening, not to label individuals.",
            "Pair satisfaction data with manager coaching and work-life-balance interventions.",
            "Track model calibration and intervention outcomes over time; do not optimize only for accuracy.",
        ],
    )

    add_heading(document, "10. Conclusion")
    document.add_paragraph(
        "This project converts a static IBM HR snapshot into a reproducible retention analytics workflow. "
        "It demonstrates data hygiene, structured EDA, leakage-aware modeling, explainable scoring, and a "
        "usable dashboard while keeping human judgment and fairness guardrails explicit."
    )

    path = REPORTS / "IBM_HR_Attrition_ProjectReport.docx"
    document.save(path)
    print(path)


if __name__ == "__main__":
    main()